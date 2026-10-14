# -*- coding: utf-8 -*-
"""从原稿提取 12 张图（按图号命名）并导出 4 张表为 CSV 到 JTM_submission/"""
import csv, os, re, sys
try:
    from docx import Document
except ImportError:
    sys.path.insert(0, r'C:/Users/WZ/.workbuddy/binaries/python/envs/default/Lib/site-packages')
    from docx import Document
from docx.oxml.ns import qn

W4D = r"D:/文章/Kimi_Agent_DFSP单细胞研究/DFSP/analysis/W4_写作"
SRC = os.path.join(W4D, "manuscript_DFSP_EN.docx")
SUB = os.path.join(W4D, "JTM_submission")
FIG_DIR = os.path.join(SUB, "figures")
TBL_DIR = os.path.join(SUB, "tables")
os.makedirs(FIG_DIR, exist_ok=True)
os.makedirs(TBL_DIR, exist_ok=True)

doc = Document(SRC)

# ============ 1) 提取图：按 body 顺序，图片段后紧跟 "Figure N" 图注 ============
from docx.text.paragraph import Paragraph
body = doc.element.body
paras = doc.paragraphs

def iter_body_paras():
    for p in body.iterchildren():
        if p.tag == qn('w:p'):
            yield Paragraph(p, doc)

def para_has_image(p):
    return len(p._element.findall('.//' + qn('w:drawing'))) > 0 or 'graphic' in p._element.xml

def extract_image(p):
    """返回 (扩展名, bytes)"""
    for blip in p._element.findall('.//' + qn('a:blip')):
        rid = blip.get(qn('r:embed'))
        if rid:
            part = doc.part.related_parts[rid]
            ext = part.partname.ext or 'png'
            return ext, part.blob
    return None, None

# 顺序遍历：图段 + 下一个非空图注段配对
figs = []  # (图号, ext, blob)
pending_img = None
for p in iter_body_paras():
    if para_has_image(p):
        ext, blob = extract_image(p)
        pending_img = (ext, blob)
    else:
        t = p.text.strip()
        m = re.match(r"^(Figure \d+[A-C]?)\.", t)
        if m and pending_img:
            figs.append((m.group(1).replace(" ", ""), pending_img[0], pending_img[1]))
            pending_img = None
        elif t:
            pending_img = None

print(f"提取图数: {len(figs)}")
for label, ext, blob in figs:
    fname = f"{label}.{ext}"
    with open(os.path.join(FIG_DIR, fname), "wb") as f:
        f.write(blob)
    print(f"  {fname} ({len(blob)//1024} KB)")

# ============ 2) 导出表格 CSV ============
table_names = ["Table1_datasets", "SupplementaryTable1_151gene_signature",
               "SupplementaryTable2_panel_genes", "SupplementaryTable3_EGFR_drugs"]
for i, tbl in enumerate(doc.tables):
    fname = os.path.join(TBL_DIR, table_names[i] + ".csv")
    with open(fname, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        for row in tbl.rows:
            w.writerow([c.text.strip() for c in row.cells])
    print(f"  表导出: {table_names[i]}.csv ({len(tbl.rows)} 行)")

# ============ 3) README 说明 ============
readme = """# JTM Submission 独立文件目录

## figures/（12 张图，按图号命名，与原稿内嵌一致）
- Figure1.* ~ Figure9.*；其中 Figure2A/2B、Figure3A/3B、Figure6A/6B 为同一图的两个面板
- 上传前建议将 Figure2A+2B、Figure3A+3B、Figure6A+6B 各合并为一个复合文件（面板 a/b）
- 分辨率要求: 300 dpi @ 最终尺寸；单文件 ≤ 10 MB；图内键（legend key）嵌入图内

## tables/
- Table1_datasets.csv       — 主表 Table 1（公共数据集）
- SupplementaryTable1_151gene_signature.csv — 151 基因共识签名（152 行，建议作为附加文件上传）
- SupplementaryTable2_panel_genes.csv       — ML 诊断 panel 基因
- SupplementaryTable3_EGFR_drugs.csv        — EGFR 轴候选药物（ChEMBL）

## 上传建议（JTM 投稿系统）
1. 主文稿: manuscript_DFSP_JTM_ready.docx（已移除内嵌图，图注集中在 References 后）
2. 图: figures/ 下 12 个文件（或合并后的 9 个复合文件）
3. 表: Table1 已含在主文稿文末；大表 SupplementaryTable1 建议另存为附加文件（.csv/.xlsx）
4. 附加文件描述已写入主文稿末尾（Additional files 段）

"""
with open(os.path.join(SUB, "README.md"), "w", encoding="utf-8") as f:
    f.write(readme)
print("\nJTM_submission 目录就绪:", SUB)
